"""Merge the UzNER-Style v2 workbook into the project's NER JSONL dataset."""

import argparse
import hashlib
import json
from collections import Counter
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

try:
    from scripts.merge_mendeley_ner import (
        prepare_output_dir,
        read_jsonl,
        validate_target_entities,
        write_jsonl,
    )
except ModuleNotFoundError:
    from merge_mendeley_ner import (
        prepare_output_dir,
        read_jsonl,
        validate_target_entities,
        write_jsonl,
    )


TARGET_LABELS = ("ORG", "NAME", "GEO")
SOURCE_TO_TARGET = {
    "PER": "NAME",
    "PERSON": "NAME",
    "LOC": "GEO",
    "LOCATION": "GEO",
    "GPE": "GEO",
    "GEOPOLITICALENTITY": "GEO",
    "ORG": "ORG",
    "ORGANIZATION": "ORG",
}
DEFAULT_XLSX = Path("useful_files/UzNER-Style(v.2).xlsx")
DEFAULT_BASE_TRAIN = Path("data/merged_mendeley_v5/train.jsonl")
DEFAULT_BASE_DEV = Path("data/merged_mendeley_v5/dev.jsonl")
DEFAULT_OUTPUT_DIR = Path("data/merged_mendeley_uzner_v2")
DEFAULT_SHEETS = ("Dataset_1", "Dataset_2")
NO_SPACE_BEFORE = frozenset(",.!?;:%)]}" + "»”’")
NO_SPACE_AFTER = frozenset("([{«“")
JsonObject = dict[str, Any]


def parse_args() -> argparse.Namespace:
    """Parse workbook merge options."""

    parser = argparse.ArgumentParser(
        description="Merge UzNER-Style v2 BIOES annotations into project JSONL."
    )
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--base-train", type=Path, default=DEFAULT_BASE_TRAIN)
    parser.add_argument("--base-dev", type=Path, default=DEFAULT_BASE_DEV)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument(
        "--sheet",
        dest="sheets",
        action="append",
        help="Workbook sheet to import; may be repeated. Defaults to Dataset_1 and Dataset_2.",
    )
    parser.add_argument("--max-sentences", type=int)
    parser.add_argument("--overwrite", action="store_true")
    return parser.parse_args()


def _column_value(row: tuple[Any, ...], index: int, source: str) -> str:
    """Read one required workbook cell as a non-empty string."""

    if index >= len(row) or row[index] is None:
        raise ValueError(f"{source}: missing column {index + 1}")
    value = str(row[index]).strip()
    if not value:
        raise ValueError(f"{source}: empty column {index + 1}")
    return value


def _normalise_label(label: str) -> str:
    """Normalize source label spelling for the mapping table."""

    return "".join(character for character in label.upper() if character.isalnum())


def _parse_tag(raw_tag: str, source: str) -> tuple[str, str | None]:
    """Parse a BIOES tag and return its prefix and normalized source label."""

    tag = raw_tag.upper().replace("_", "-")
    if tag == "O":
        return "O", None
    prefix, separator, label = tag.partition("-")
    if not separator or prefix not in {"B", "I", "O", "E", "S"} or not label:
        raise ValueError(f"{source}: unsupported BIOES tag {raw_tag!r}")
    return prefix, _normalise_label(label)


def _append_word(words: list[str], word: str) -> tuple[int, int]:
    """Append a token with a readable detokenization and return its character span."""

    if not words:
        words.append(word)
        return 0, len(word)
    previous = words[-1]
    separator = "" if word[0] in NO_SPACE_BEFORE or previous[-1] in NO_SPACE_AFTER else " "
    start = sum(len(item) for item in words) + len(separator)
    words.append(separator + word)
    return start, start + len(word)


def _entities_from_tokens(
    token_spans: list[tuple[int, int]],
    raw_tags: list[str],
    source: str,
    counters: Counter[str],
) -> list[JsonObject]:
    """Convert token-level BIOES tags into non-overlapping target spans."""

    entities: list[JsonObject] = []
    current: tuple[str, int] | None = None

    def flush(end_token: int) -> None:
        nonlocal current
        if current is None:
            return
        source_label, start_token = current
        counters[f"source_entities:{source_label}"] += 1
        target_label = SOURCE_TO_TARGET.get(source_label)
        if target_label is None:
            counters[f"dropped_source_entities:{source_label}"] += 1
        else:
            entities.append(
                {
                    "label": target_label,
                    "start": token_spans[start_token][0],
                    "end": token_spans[end_token][1],
                }
            )
        current = None

    for index, raw_tag in enumerate(raw_tags):
        prefix, label = _parse_tag(raw_tag, f"{source}/token[{index}]")
        if prefix == "O":
            flush(index - 1)
        elif prefix == "S":
            flush(index - 1)
            current = (label or "", index)
            flush(index)
        elif prefix == "B":
            flush(index - 1)
            current = (label or "", index)
        elif prefix == "I":
            if current is None:
                counters["malformed_i_without_b"] += 1
                current = (label or "", index)
            elif current[0] != label:
                counters["malformed_label_transition"] += 1
                flush(index - 1)
                current = (label or "", index)
        elif prefix == "E":
            if current is None:
                counters["malformed_e_without_b"] += 1
                current = (label or "", index)
            elif current[0] != label:
                counters["malformed_label_transition"] += 1
                flush(index - 1)
                current = (label or "", index)
            flush(index)
    flush(len(raw_tags) - 1)
    return sorted(entities, key=lambda item: (item["start"], item["end"], item["label"]))


def _convert_sentence(
    sheet: str,
    sentence_id: str,
    rows: list[tuple[Any, ...]],
    counters: Counter[str],
) -> JsonObject:
    """Convert one grouped workbook sentence into the project schema."""

    words: list[str] = []
    spans: list[tuple[int, int]] = []
    tags: list[str] = []
    for row_index, row in enumerate(rows, start=2):
        source = f"{sheet}:{row_index}"
        word = _column_value(row, 4, source)
        tag = _column_value(row, 5, source)
        start, end = _append_word(words, word)
        spans.append((start, end))
        tags.append(tag)
    text = "".join(words)
    entities = _entities_from_tokens(
        spans,
        tags,
        f"{sheet}:{sentence_id}",
        counters,
    )
    validate_target_entities(entities, len(text), f"{sheet}:{sentence_id}")
    digest = hashlib.sha256(f"{sheet}\0{sentence_id}\0{text}".encode()).hexdigest()
    return {
        "hash": f"uzner-v2-{digest}",
        "text": text,
        "entities": entities,
    }


def _sentence_groups(
    rows: Iterable[tuple[Any, ...]],
    sheet: str,
    counters: Counter[str],
    max_sentences: int | None,
) -> list[JsonObject]:
    """Group contiguous workbook token rows by sentence identifier."""

    result: list[JsonObject] = []
    current_id: str | None = None
    current_rows: list[tuple[Any, ...]] = []
    for row_number, row in enumerate(rows, start=2):
        sentence_id = _column_value(row, 3, f"{sheet}:{row_number}")
        if current_id is not None and sentence_id != current_id:
            result.append(_convert_sentence(sheet, current_id, current_rows, counters))
            if max_sentences is not None and len(result) >= max_sentences:
                return result
            current_rows = []
        current_id = sentence_id
        current_rows.append(row)
    if current_rows and (max_sentences is None or len(result) < max_sentences):
        result.append(_convert_sentence(sheet, current_id or "unknown", current_rows, counters))
    return result


def load_uzner_records(
    path: Path,
    sheets: tuple[str, ...],
    max_sentences: int | None = None,
) -> tuple[list[JsonObject], Counter[str]]:
    """Stream selected workbook sheets and convert their annotations."""

    counters: Counter[str] = Counter()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        missing = [sheet for sheet in sheets if sheet not in workbook.sheetnames]
        if missing:
            raise ValueError(f"{path}: missing sheets {missing}; available={workbook.sheetnames}")
        records: list[JsonObject] = []
        for sheet_name in sheets:
            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(min_row=2, values_only=True)
            sheet_records = _sentence_groups(rows, sheet_name, counters, max_sentences)
            records.extend(sheet_records)
            counters[f"sentences:{sheet_name}"] = len(sheet_records)
    finally:
        workbook.close()
    if not records:
        raise ValueError(f"{path}: no sentences were converted")
    return records, counters


def build_manifest(
    xlsx_path: Path,
    base_train: Path,
    base_dev: Path,
    output_dir: Path,
    sheets: tuple[str, ...],
    original_train: list[JsonObject],
    original_dev: list[JsonObject],
    added_records: list[JsonObject],
    counters: Counter[str],
) -> JsonObject:
    """Build an auditable merge manifest."""

    entity_counts = Counter(
        entity["label"] for record in added_records for entity in record["entities"]
    )
    return {
        "schema_version": 1,
        "operation": "merge_uzner_style_v2_into_train",
        "source_file": str(xlsx_path),
        "source_url": "https://data.mendeley.com/datasets/48923w3gyr/1",
        "source_license": "CC BY 4.0",
        "sheets": list(sheets),
        "base_train": str(base_train),
        "base_dev": str(base_dev),
        "policy": "append converted Dataset_1 and Dataset_2 records to train; keep dev unchanged",
        "label_mapping": SOURCE_TO_TARGET,
        "tag_scheme": "BIOES converted to exact character spans",
        "counts": {
            "original_train_records": len(original_train),
            "original_dev_records": len(original_dev),
            "added_records": len(added_records),
            "merged_train_records": len(original_train) + len(added_records),
            "merged_dev_records": len(original_dev),
            "added_entities_by_target_label": dict(sorted(entity_counts.items())),
            **dict(sorted(counters.items())),
        },
        "outputs": {
            "train": str(output_dir / "train.jsonl"),
            "dev": str(output_dir / "dev.jsonl"),
            "manifest": str(output_dir / "merge_manifest.json"),
        },
    }


def run(args: argparse.Namespace) -> int:
    """Convert the workbook and merge it into the existing training split."""

    xlsx_path = args.xlsx.expanduser().resolve()
    base_train = args.base_train.expanduser().resolve()
    base_dev = args.base_dev.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    for path in (xlsx_path, base_train, base_dev):
        if not path.is_file():
            raise FileNotFoundError(path)

    sheets = tuple(args.sheets or DEFAULT_SHEETS)
    original_train = read_jsonl(base_train)
    original_dev = read_jsonl(base_dev)
    external_records, counters = load_uzner_records(
        xlsx_path,
        sheets,
        max_sentences=args.max_sentences,
    )
    existing_texts = {record["text"] for record in original_train + original_dev}
    seen_texts = set(existing_texts)
    added_records: list[JsonObject] = []
    for record in external_records:
        if record["text"] in seen_texts:
            counters["skipped_duplicate_text"] += 1
            continue
        seen_texts.add(record["text"])
        added_records.append(record)

    prepare_output_dir(output_dir, args.overwrite)
    write_jsonl(output_dir / "train.jsonl", original_train + added_records)
    write_jsonl(output_dir / "dev.jsonl", original_dev)
    manifest = build_manifest(
        xlsx_path,
        base_train,
        base_dev,
        output_dir,
        sheets,
        original_train,
        original_dev,
        added_records,
        counters,
    )
    (output_dir / "merge_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"Original train: {len(original_train)}")
    print(f"Original dev: {len(original_dev)}")
    print(f"Added UzNER records: {len(added_records)}")
    print(f"Merged train: {len(original_train) + len(added_records)}")
    print(f"Dev remains unchanged: {len(original_dev)}")
    print(f"Output train: {output_dir / 'train.jsonl'}")
    print(f"Manifest: {output_dir / 'merge_manifest.json'}")
    print(f"Target entities: {manifest['counts']['added_entities_by_target_label']}")
    return 0


def main() -> int:
    """Run the merge and report validation errors."""

    try:
        return run(parse_args())
    except (OSError, TypeError, ValueError, json.JSONDecodeError) as error:
        print(f"ERROR: {error}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
